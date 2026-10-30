from collections import defaultdict
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from database import list_candidates
from config import EXPORT_FILE

VALID_ANALYSIS_STATUSES = ("VALIDO",)


def _parse_data_br(data: str):
    try:
        return datetime.strptime(data, "%d/%m/%Y")
    except Exception:
        return datetime.min


def whatsapp_display(valor):
    return valor if valor else "Número não informado"


def _parse_data_hora_br(data: str | None, hora: str | None):
    try:
        return datetime.strptime(f"{data} {hora}", "%d/%m/%Y %H:%M:%S")
    except Exception:
        return datetime.min


def marcar_pesquisas_incompletas(conn, minutos: int = 30) -> None:
    """Marca como INCOMPLETA pesquisas abandonadas há mais de X minutos.

    Os duelos já registrados continuam salvos e aparecem no painel.
    Pesquisas recentes em andamento continuam com status EM_ANDAMENTO.
    """
    limite = datetime.now() - timedelta(minutes=minutos)
    rows = conn.execute("SELECT * FROM pesquisas WHERE status = 'EM_ANDAMENTO'").fetchall()
    for r in rows:
        ultimo_ts = r["ultimo_voto_ts"]
        if ultimo_ts:
            ultima = datetime.fromtimestamp(float(ultimo_ts))
        else:
            ultima = _parse_data_hora_br(r["data_inicio"], r["hora_inicio"])
        if ultima and ultima < limite:
            conn.execute(
                """
                UPDATE pesquisas
                SET status = 'INCOMPLETA',
                    motivo_status = 'Pesquisa abandonada: usuário saiu, fechou a página ou ficou sem atividade por mais de 30 minutos.',
                    ganhador = COALESCE(current_left, ganhador),
                    data_fim = COALESCE(data_fim, data_inicio),
                    hora_fim = COALESCE(hora_fim, hora_inicio)
                WHERE id = ?
                """,
                (r["id"],),
            )


def dashboard_context(conn):
    marcar_pesquisas_incompletas(conn)
    total_validas = conn.execute("SELECT COUNT(*) AS t FROM pesquisas WHERE status = 'VALIDO'").fetchone()["t"]
    total_pesquisas = conn.execute("SELECT COUNT(*) AS t FROM pesquisas").fetchone()["t"]
    total_duelos = conn.execute("SELECT COUNT(*) AS t FROM duelos").fetchone()["t"]
    participantes_unicos = conn.execute(
        """
        SELECT COUNT(*) AS t FROM (
          SELECT COALESCE(NULLIF(whatsapp,''), NULLIF(fingerprint,''), ip) AS chave
          FROM pesquisas
          GROUP BY chave
        )
        """
    ).fetchone()["t"]
    ranking_completo = ranking_stats(conn)
    lider = ranking_completo[0] if ranking_completo else None
    pesquisas = [dict(r) for r in conn.execute("SELECT * FROM pesquisas WHERE status IN ('VALIDO','SUSPEITO','IGNORADO_REPETICAO') ORDER BY id DESC LIMIT 500").fetchall()]
    pesquisas_incompletas = [dict(r) for r in conn.execute("SELECT * FROM pesquisas WHERE status IN ('EM_ANDAMENTO','INCOMPLETA') ORDER BY id DESC LIMIT 500").fetchall()]
    auditoria = [dict(r) for r in conn.execute("SELECT * FROM pesquisas WHERE status IN ('SUSPEITO','IGNORADO_REPETICAO','INCOMPLETA') ORDER BY id DESC LIMIT 500").fetchall()]
    evolucao = evolution_by_day(conn)
    status_chart = status_counts(conn)
    return {
        "total_validas": total_validas,
        "total_pesquisas": total_pesquisas,
        "participantes_unicos": participantes_unicos,
        "total_duelos": total_duelos,
        "lider": lider,
        "ranking_completo": ranking_completo,
        "top10": ranking_completo[:10],
        "pesquisas": pesquisas,
        "pesquisas_incompletas": pesquisas_incompletas,
        "auditoria": auditoria,
        "evolucao": evolucao,
        "status_chart": status_chart,
        "chart_payload": {
            "evolucao": evolucao,
            "top10": ranking_completo[:10],
            "status": status_chart,
        },
    }


def status_counts(conn):
    rows = conn.execute("SELECT status, COUNT(*) AS total FROM pesquisas GROUP BY status").fetchall()
    return [{"status": r["status"], "total": r["total"]} for r in rows]


def evolution_by_day(conn):
    rows = conn.execute(
        """
        SELECT data_fim AS data, COUNT(*) AS total
        FROM pesquisas
        WHERE status = 'VALIDO' AND data_fim IS NOT NULL
        GROUP BY data_fim
        """
    ).fetchall()
    data = sorted([dict(r) for r in rows], key=lambda x: _parse_data_br(x["data"]))
    return data


def ranking_stats(conn):
    candidatos = [dict(r) for r in list_candidates(conn, only_active=False)]
    pesquisas_rows = conn.execute("SELECT ganhador, COUNT(*) AS total FROM pesquisas WHERE status = 'VALIDO' AND ganhador IS NOT NULL GROUP BY ganhador").fetchall()
    pesquisas_ganhas = {r["ganhador"]: r["total"] for r in pesquisas_rows}

    duelos_v = defaultdict(int)
    duelos_p = defaultdict(int)
    for r in conn.execute("SELECT candidato_votado, candidato_adversario FROM duelos").fetchall():
        duelos_v[r["candidato_votado"]] += 1
        duelos_p[r["candidato_adversario"]] += 1

    total_validas = sum(pesquisas_ganhas.values()) or 1
    lista = []
    for c in candidatos:
        nome = c["nome"]
        v = duelos_v[nome]
        p = duelos_p[nome]
        total_confrontos = v + p
        taxa = (v / total_confrontos * 100) if total_confrontos else 0
        ganhos = pesquisas_ganhas.get(nome, 0)
        lista.append({
            "vereador": nome,
            "partido": c["partido"],
            "img": c["img"],
            "pesquisas_ganhas": ganhos,
            "percentual": (ganhos / total_validas * 100) if total_validas else 0,
            "duelos_vencidos": v,
            "duelos_perdidos": p,
            "total_confrontos": total_confrontos,
            "taxa_vitoria": taxa,
        })
    lista.sort(key=lambda x: (x["pesquisas_ganhas"], x["duelos_vencidos"], x["taxa_vitoria"]), reverse=True)
    for i, item in enumerate(lista, start=1):
        item["posicao"] = i
    return lista


def pesquisa_detalhes(conn, pesquisa_id: int):
    p = conn.execute("SELECT * FROM pesquisas WHERE id = ?", (pesquisa_id,)).fetchone()
    if not p:
        return None
    duelos = conn.execute("SELECT * FROM duelos WHERE pesquisa_id = ? ORDER BY id", (pesquisa_id,)).fetchall()
    return {"pesquisa": dict(p), "duelos": [dict(d) for d in duelos]}


def _style_header(ws):
    fill = PatternFill("solid", fgColor="14213D")
    font = Font(bold=True, color="FFFFFF")
    border = Border(left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"), top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"))
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    ws.freeze_panes = "A2"


def gerar_excel(conn):
    marcar_pesquisas_incompletas(conn)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pesquisas"
    ws.append(["ID", "Data Início", "Hora Início", "Data Fim", "Hora Fim", "IP", "WhatsApp", "Fingerprint", "Ganhador", "Total Duelos", "Status", "Motivo"])
    for r in conn.execute("SELECT * FROM pesquisas ORDER BY id").fetchall():
        ws.append([r["id"], r["data_inicio"], r["hora_inicio"], r["data_fim"], r["hora_fim"], r["ip"], whatsapp_display(r["whatsapp"]), r["fingerprint"], r["ganhador"], r["total_duelos"], r["status"], r["motivo_status"]])
    _style_header(ws)

    ws = wb.create_sheet("Incompletas")
    ws.append(["ID", "Data Início", "Hora Início", "Última Atividade", "IP", "WhatsApp", "Fingerprint", "Ganhador Parcial", "Total Duelos", "Status", "Motivo"])
    for r in conn.execute("SELECT * FROM pesquisas WHERE status IN ('EM_ANDAMENTO','INCOMPLETA') ORDER BY id").fetchall():
        ultima = datetime.fromtimestamp(float(r["ultimo_voto_ts"])).strftime("%d/%m/%Y %H:%M:%S") if r["ultimo_voto_ts"] else f"{r['data_inicio']} {r['hora_inicio']}"
        ws.append([r["id"], r["data_inicio"], r["hora_inicio"], ultima, r["ip"], whatsapp_display(r["whatsapp"]), r["fingerprint"], r["current_left"] or r["ganhador"], r["total_duelos"], r["status"], r["motivo_status"]])
    _style_header(ws)

    ws = wb.create_sheet("Ranking")
    ws.append(["Posição", "Vereador", "Partido", "Pesquisas Ganhas", "% Preferência", "Duelos Vencidos", "Duelos Perdidos", "Total Confrontos", "Taxa Vitória"])
    for r in ranking_stats(conn):
        ws.append([r["posicao"], r["vereador"], r["partido"], r["pesquisas_ganhas"], round(r["percentual"], 2), r["duelos_vencidos"], r["duelos_perdidos"], r["total_confrontos"], round(r["taxa_vitoria"], 2)])
    _style_header(ws)

    ws = wb.create_sheet("Duelos")
    ws.append(["ID", "Pesquisa", "Data", "Horário", "IP", "Vencedor", "Adversário", "Chave", "Fingerprint"])
    for r in conn.execute("SELECT * FROM duelos ORDER BY id").fetchall():
        ws.append([r["id"], r["pesquisa_id"], r["data"], r["horario"], r["ip"], r["candidato_votado"], r["candidato_adversario"], r["chave_duelo"], r["fingerprint"]])
    _style_header(ws)

    ws = wb.create_sheet("Auditoria")
    ws.append(["ID", "Data Fim", "IP", "WhatsApp", "Fingerprint", "Ganhador", "Status", "Motivo"])
    for r in conn.execute("SELECT * FROM pesquisas WHERE status IN ('SUSPEITO','IGNORADO_REPETICAO') ORDER BY id").fetchall():
        ws.append([r["id"], r["data_fim"], r["ip"], whatsapp_display(r["whatsapp"]), r["fingerprint"], r["ganhador"], r["status"], r["motivo_status"]])
    _style_header(ws)

    ws = wb.create_sheet("Evolução")
    ws.append(["Data", "Pesquisas Válidas"])
    for r in evolution_by_day(conn):
        ws.append([r["data"], r["total"]])
    _style_header(ws)

    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 2, 45)
    wb.save(EXPORT_FILE)
    return EXPORT_FILE
